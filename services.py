"""Vietnam product services. CV input is never persisted or sent over the network."""
import io
import json
import re
import zipfile
from copy import deepcopy
from functools import lru_cache
from datetime import datetime, timezone
from urllib.parse import urlsplit
from intelligence import enrich, plain, SKILLS, skill_tags, visible_jobs
from locations import location_tags, folded

UNKNOWN = 'Chưa công bố'
FIELDS = {'greenhouse': 'board_token', 'lever': 'site', 'ashby': 'board_name', 'smartrecruiters': 'company_identifier'}


def safe_url(value):
    try:
        u = urlsplit(str(value or ''))
        return u.scheme == 'https' and bool(u.hostname) and not u.username and not u.password
    except ValueError:
        return False


@lru_cache(maxsize=4096)
def _enriched_record(serialized):
    return enrich(json.loads(serialized))


def vietnam_jobs(records, active_only=True):
    result = []
    for record in records:
        if not isinstance(record, dict) or not record.get('id') or not record.get('title'):
            continue
        job = deepcopy(_enriched_record(json.dumps(record, ensure_ascii=False, sort_keys=True)))
        if job.get('sections'):
            job['skills'] = skill_tags(matching_text(job))
        job['source_location'] = record.get('source_location') or record.get('location', '')
        job['locations'] = location_tags(job['source_location'])
        job['location'] = ' / '.join(job['locations'])
        if display(job.get('experience')) == UNKNOWN:
            found = re.search(r'\b\d+(?:\s*[-–]\s*\d+)?\+?\s*(?:years?(?: of)? experience|năm kinh nghiệm)', job['description'], re.I)
            if found:
                job['experience'] = found.group(0)
        if job['country'] != 'Vietnam' or not safe_url(job.get('url')):
            continue
        if active_only and job.get('status', 'active') != 'active':
            continue
        result.append(job)
    return visible_jobs(result)


def display(value):
    return UNKNOWN if not value or value in ('Not specified', 'Unknown', 'Not disclosed') else str(value)


def sentences(text):
    text = re.sub(r'</(?:p|li|div)>|<br\s*/?>', '\n', str(text or ''), flags=re.I)
    return [plain(s) for s in re.split(r'\n|(?<=[.!?;])\s+', text) if plain(s)]


def dossier(job):
    lines = sentences(job.get('description', ''))
    groups = {
        'Ưu tiên': r'preferred|nice to have|advantage|ưu tiên|lợi thế',
        'Kinh nghiệm': r'experience|years|kinh nghiệm|\d+\s*năm',
        'Học vấn & chứng chỉ': r'degree|bachelor|master|certifi|đại học|bằng cấp|chứng chỉ',
        'Ngôn ngữ': r'english|vietnamese|japanese|chinese|tiếng anh|tiếng nhật|tiếng trung',
        'Kỹ năng & yêu cầu khác': r'requir|must|proficien|skill|yêu cầu|thành thạo|kỹ năng',
    }
    requirements = {label: [] for label in groups}
    duties = []
    for line in lines:
        for label, pattern in groups.items():
            if re.search(pattern, line, re.I):
                requirements[label].append(line)
                break
        else:
            if re.search(r'manage|develop|deliver|support|responsib|analy[sz]|build|coordinate|quản lý|thực hiện|phân tích|hỗ trợ|phối hợp|xây dựng', line, re.I):
                duties.append(line)
    return {'overview': lines[0] if lines else '', 'duties': duties[:8], 'requirements': requirements}


def matching_text(job):
    if job.get('sections'):
        return '\n'.join(s['text'] for s in job['sections'] if s.get('title') != 'Về doanh nghiệp')
    return job.get('description', '')


def match_cv(cv, job):
    """Literal evidence is partial; never claim competence from a keyword."""
    report = []
    for skill in job.get('skills', []):
        pattern = SKILLS[skill]
        jd = next((s for s in sentences(matching_text(job)) if re.search(pattern, s, re.I)), '')
        evidence = next((s for s in sentences(cv) if re.search(pattern, s, re.I)), '')
        applied = bool(evidence and re.search(r'built|developed|delivered|managed|used|implemented|created|analyzed|xây dựng|sử dụng|triển khai|phân tích|quản lý', evidence, re.I))
        report.append({'Yêu cầu': skill, 'Phân loại': 'Đáp ứng' if applied else 'Đáp ứng một phần' if evidence else 'Chưa tìm thấy',
                       'Dẫn chứng JD': jd, 'Dẫn chứng CV': evidence or 'Không tìm thấy trong văn bản đã trích xuất.',
                       'Gợi ý': 'Bổ sung kết quả và phạm vi dự án để xác minh mức thành thạo.' if evidence else 'Bổ sung trải nghiệm liên quan nếu có; không khai kỹ năng chưa sở hữu.'})
    return report


def rank_jobs(cv, jobs, query=''):
    """Transparent lexical relevance, not hiring probability. No external CV calls."""
    import math
    from collections import Counter
    stop = set('the and for with that this from your you our are will have has can all job work company team experience skills requirements vietnam trong cong viec va cua voi cac cho mot duoc'.split())
    def tokens(text):
        return {t for t in re.findall(r'[a-z][a-z0-9+#.]*', folded(text)) if len(t) > 2 and t not in stop}
    candidates = [j for j in jobs if all(word in folded(' '.join(str(j.get(k, '')) for k in ('title','company','location','description','skills'))) for word in folded(query).split())]
    docs = [tokens(j.get('title', '') + ' ' + matching_text(j)) for j in jobs]
    frequency = Counter(t for doc in docs for t in doc)
    weights = {t: math.log(1 + len(jobs) / n) for t, n in frequency.items()}
    cv_tokens = tokens(cv)
    output = []
    for job in candidates:
        report = match_cv(cv, job) if cv else []
        matched = [r['Yêu cầu'] for r in report if r['Phân loại'] != 'Chưa tìm thấy']
        terms = tokens(job.get('title', '') + ' ' + matching_text(job))
        overlap = terms & cv_tokens
        lexical = sum(weights.get(t, 0) for t in overlap) / max(1, sum(weights.get(t, 0) for t in terms))
        skill_ratio = len(matched) / len(report) if report else lexical
        score = round(100 * (.65 * skill_ratio + .35 * lexical), 1) if cv and job.get('description') else None
        output.append({'job': job, 'score': score, 'matched': matched, 'required': len(report), 'keywords': sorted(overlap, key=lambda t: (-weights.get(t, 0), t))[:8]})
    return sorted(output, key=lambda r: (r['score'] is None, -(r['score'] or 0), -len(r['matched']), r['job']['title'].casefold(), r['job']['id']))


def parse_cv(data, filename):
    if len(data) > 5 * 1024 * 1024:
        raise ValueError('CV tối đa 5 MB.')
    if filename.lower().endswith('.pdf'):
        from pypdf import PdfReader
        reader = PdfReader(io.BytesIO(data))
        if reader.is_encrypted or len(reader.pages) > 30:
            raise ValueError('Dùng PDF không mã hóa, tối đa 30 trang.')
        text = '\n'.join(p.extract_text() or '' for p in reader.pages)
    elif filename.lower().endswith('.docx'):
        from docx import Document
        with zipfile.ZipFile(io.BytesIO(data)) as archive:
            if sum(z.file_size for z in archive.infolist()) > 20 * 1024 * 1024:
                raise ValueError('Tệp DOCX giải nén vượt giới hạn 20 MB.')
        document = Document(io.BytesIO(data))
        text = '\n'.join([p.text for p in document.paragraphs] + [c.text for t in document.tables for r in t.rows for c in r.cells])
    else:
        raise ValueError('Chỉ nhận PDF hoặc DOCX.')
    if not text.strip():
        raise ValueError('Không đọc được văn bản. PDF dạng ảnh cần chuyển thành văn bản trước.')
    return text[:100000]


def detect_board(url):
    if not safe_url(url):
        raise ValueError('Cần URL HTTPS hợp lệ.')
    u = urlsplit(url)
    host = u.hostname.lower()
    parts = [x for x in u.path.split('/') if x]
    hosts = {'boards.greenhouse.io': 'greenhouse', 'job-boards.greenhouse.io': 'greenhouse', 'jobs.lever.co': 'lever', 'jobs.ashbyhq.com': 'ashby', 'careers.smartrecruiters.com': 'smartrecruiters', 'jobs.smartrecruiters.com': 'smartrecruiters'}
    provider = hosts.get(host)
    if not provider or not parts or not re.fullmatch(r'[A-Za-z0-9_-]+', parts[0]):
        raise ValueError('Chưa nhận diện được ATS. Lưu vào danh sách doanh nghiệp để xác minh thủ công.')
    return provider, parts[0]


def register_source(config, company, url, permission='', enabled=False, permission_mode='full', ai_authorized=False):
    provider, board = detect_board(url)
    if not company.strip():
        raise ValueError('Cần tên doanh nghiệp.')
    if enabled and not safe_url(permission):
        raise ValueError('Cần liên kết bằng chứng cho phép sử dụng dữ liệu trước khi kích hoạt.')
    if permission_mode not in ('full', 'links'):
        raise ValueError('Phạm vi sử dụng dữ liệu không hợp lệ.')
    updated = json.loads(json.dumps(config))
    entries = updated.setdefault(provider, [])
    old = next((s for s in entries if s.get(FIELDS[provider]) == board), None)
    item = {**(old or {}), 'company': company.strip(), FIELDS[provider]: board, 'careers_url': url,
            'enabled': enabled, 'authorized': enabled and permission_mode == 'full', 'discovery_enabled': enabled and permission_mode == 'links',
            'permission_mode': permission_mode, 'ai_authorized': enabled and permission_mode == 'full' and ai_authorized,
            'permission_url': permission, 'permission_checked_at': datetime.now(timezone.utc).isoformat() if enabled else None}
    if old is None:
        entries.append(item)
    else:
        entries[entries.index(old)] = item
    return updated


CAREERS = {
    'Phân tích dữ liệu': ['Data Analyst', 'Senior Data Analyst', 'Analytics Manager', 'Head of Analytics'],
    'Vận hành & bán lẻ': ['Store Manager', 'Operations Excellence', 'Business Operations', 'Strategy & Operations'],
    'Kỹ thuật phần mềm': ['Software Engineer', 'Senior Software Engineer', 'Tech Lead', 'Engineering Manager'],
    'Tài chính': ['Financial Analyst', 'Senior Financial Analyst', 'Finance Manager', 'Finance Director'],
    'Marketing': ['Marketing Executive', 'Marketing Specialist', 'Marketing Manager', 'Head of Marketing'],
    'Sản xuất': ['Production Engineer', 'Production Supervisor', 'Production Manager', 'Plant Manager'],
    'Nhân sự': ['HR Executive', 'HR Specialist', 'HR Business Partner', 'HR Manager'],
    'Y tế': ['Điều dưỡng', 'Điều dưỡng trưởng', 'Quản lý điều dưỡng'],
}
